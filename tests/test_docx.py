"""DOCX structure: headings from Word styles/numbering, notes, headers. No DB or API needed."""
from docx import Document

from compliance_rag.chunking.splitter import chunk_doc
from compliance_rag.ingestion.parsers import note_lines, parse_docx


def make_contract(path):
    d = Document()
    d.sections[0].header.paragraphs[0].text = "CONFIDENTIAL - Ref SA-2026-14"
    d.add_paragraph("Supply Agreement", style="Title")
    d.add_paragraph("Definitions", style="Heading 1")
    d.add_paragraph("In this Agreement, Goods means the items listed in Schedule A.")
    d.add_paragraph("Termination", style="List Number")          # auto-numbered title
    d.add_paragraph("Either party may terminate on 30 days' written notice.")
    d.add_paragraph("A breach not remedied in 14 days ends this Agreement.",
                    style="List Number")                          # numbered body sentence
    d.add_paragraph().add_run("Confidentiality").bold = True      # bold title, no style
    d.add_paragraph("Each party shall keep the other's information confidential.")
    table = d.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Item", "Price"
    table.cell(1, 0).text, table.cell(1, 1).text = "Widget", "10"
    d.save(path)


def test_docx_headings_come_from_word_structure(tmp_path):
    path = tmp_path / "supply.docx"
    make_contract(path)
    doc = parse_docx(path)
    assert {"Supply Agreement", "Definitions", "Termination", "Confidentiality"} <= doc.headings
    assert "A breach not remedied in 14 days ends this Agreement." not in doc.headings
    assert "Headers and footers" in doc.headings

    by_section = {c.section: c for c in chunk_doc(doc)}
    assert "30 days' written notice" in by_section["Termination"].text
    assert "14 days" in by_section["Termination"].text      # numbered sentence stays in body
    assert by_section["Confidentiality"].chunk_type == "clause"
    assert "SA-2026-14" in by_section["Headers and footers"].text
    assert by_section["(table)"].text == "Item | Price\nWidget | 10"


def test_note_lines_skip_separators():
    w = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    xml = f"""<w:footnotes xmlns:w="{w}">
      <w:footnote w:type="separator" w:id="-1"><w:p><w:r><w:t>---</w:t></w:r></w:p></w:footnote>
      <w:footnote w:id="1"><w:p><w:r><w:t>"Affiliate" has the meaning </w:t></w:r>
        <w:r><w:t>in section 2.</w:t></w:r></w:p></w:footnote>
    </w:footnotes>""".encode()
    assert note_lines(xml, "footnote") == ['[1] "Affiliate" has the meaning in section 2.']
